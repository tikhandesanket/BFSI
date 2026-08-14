"""Persistent store for user-uploaded documents.

Each uploaded file becomes one *upload* record (uuid) and produces N *chunk*
records that feed into the RAG index. The store routes each file format to
the right parser and picks the right chunking strategy:

  * Unstructured (.txt .md .pdf .docx .html): paragraph-aware character chunks.
  * Structured (.csv .xlsx .json-array): row/record-atomic chunks where each
    record is rendered as "field: value" lines and never split mid-row.

JSON is treated as structured when it's an array of objects, otherwise
unstructured. DOCX is unstructured for prose with any tables flattened into
structured records appended to the end.
"""

from __future__ import annotations

import csv
import io
import json
import re
import time
import uuid
from dataclasses import asdict, dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from typing import List, Optional, Tuple

from .chunker import Chunk, chunk_records, chunk_text


SUPPORTED_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".pdf",
    ".json",
    ".csv",
    ".xlsx",
    ".docx",
    ".html",
    ".htm",
}

# Maps a file extension to a human-readable format type for the UI.
FORMAT_LABEL = {
    ".txt": "text",
    ".md": "markdown",
    ".markdown": "markdown",
    ".pdf": "pdf",
    ".json": "json",
    ".csv": "csv",
    ".xlsx": "excel",
    ".docx": "docx",
    ".html": "html",
    ".htm": "html",
}


@dataclass
class ChunkRecord:
    id: str              # globally unique chunk id (used as RAG doc id)
    upload_id: str       # parent upload uuid
    title: str           # human readable title (filename + chunk #)
    category: str        # category supplied at upload time (or "uploaded")
    content: str         # chunk text
    chunk_index: int
    total_chunks: int


@dataclass
class UploadRecord:
    id: str
    filename: str
    category: str
    size_bytes: int
    chunks: int
    chars: int
    uploaded_at: float
    format: str = "text"           # "text" | "pdf" | "json" | "csv" | "excel" | "docx" | "html" | "markdown"
    structure: str = "unstructured"  # "structured" | "unstructured" | "mixed"
    chunk_ids: List[str] = field(default_factory=list)

    def to_public(self) -> dict:
        return {
            "id": self.id,
            "filename": self.filename,
            "category": self.category,
            "size_bytes": self.size_bytes,
            "chunks": self.chunks,
            "chars": self.chars,
            "uploaded_at": self.uploaded_at,
            "format": self.format,
            "structure": self.structure,
        }


class UnsupportedFileType(Exception):
    pass


class EmptyDocument(Exception):
    pass


# ---------------------------------------------------------------------- helpers

def _slugify(text: str) -> str:
    text = re.sub(r"[^a-zA-Z0-9]+", "-", text).strip("-").lower()
    return text or "doc"


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise UnsupportedFileType("Could not decode file as text.")


def _stringify_cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _render_record(headers: List[str], row: List[str]) -> str:
    """Render one tabular row as `header: value` lines, skipping empty cells."""
    lines: List[str] = []
    for i, val in enumerate(row):
        v = _stringify_cell(val)
        if not v:
            continue
        h = headers[i] if i < len(headers) else f"col_{i + 1}"
        lines.append(f"{h}: {v}")
    return "\n".join(lines)


# ---------------------------------------------------------------------- parsers

def _extract_plain_text(data: bytes) -> Tuple[str, List[str]]:
    return _decode_text(data), []


def _extract_pdf(data: bytes) -> Tuple[str, List[str]]:
    try:
        from pypdf import PdfReader
    except ImportError as e:  # pragma: no cover
        raise UnsupportedFileType(
            "PDF support requires `pypdf`. Install via requirements.txt."
        ) from e
    reader = PdfReader(io.BytesIO(data))
    pages: List[str] = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001 — best effort per page
            pages.append("")
    text = "\n\n".join(p.strip() for p in pages if p.strip())
    return text, []


class _HTMLTextExtractor(HTMLParser):
    SKIP = {"script", "style", "noscript"}

    def __init__(self):
        super().__init__()
        self._buf: List[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag, _attrs):
        if tag in self.SKIP:
            self._skip_depth += 1
        elif tag in {"p", "br", "div", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"}:
            self._buf.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self._skip_depth > 0:
            self._skip_depth -= 1
        elif tag in {"p", "div", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"}:
            self._buf.append("\n")

    def handle_data(self, data):
        if self._skip_depth:
            return
        self._buf.append(data)

    def text(self) -> str:
        raw = "".join(self._buf)
        # Collapse runs of blank lines.
        lines = [ln.strip() for ln in raw.splitlines()]
        cleaned = [ln for ln in lines if ln]
        return "\n\n".join(cleaned)


def _extract_html(data: bytes) -> Tuple[str, List[str]]:
    parser = _HTMLTextExtractor()
    parser.feed(_decode_text(data))
    return parser.text(), []


def _extract_docx(data: bytes) -> Tuple[str, List[str]]:
    """DOCX: paragraphs as prose; tables flattened into structured records."""
    try:
        from docx import Document as DocxDocument
    except ImportError as e:  # pragma: no cover
        raise UnsupportedFileType(
            "DOCX support requires `python-docx`. Install via requirements.txt."
        ) from e
    doc = DocxDocument(io.BytesIO(data))
    paragraphs: List[str] = []
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            paragraphs.append(t)

    records: List[str] = []
    for tbl in doc.tables:
        rows = [
            [_stringify_cell(c.text) for c in row.cells]
            for row in tbl.rows
        ]
        if not rows:
            continue
        headers = rows[0] if any(rows[0]) else [f"col_{i + 1}" for i in range(len(rows[0]))]
        for row in rows[1:]:
            rendered = _render_record(headers, row)
            if rendered:
                records.append(rendered)

    return "\n\n".join(paragraphs), records


def _extract_csv(data: bytes) -> Tuple[str, List[str]]:
    text = _decode_text(data)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows = [list(r) for r in reader if any((cell or "").strip() for cell in r)]
    if not rows:
        return "", []
    headers = [_stringify_cell(h) or f"col_{i + 1}" for i, h in enumerate(rows[0])]
    records: List[str] = []
    for row in rows[1:]:
        rendered = _render_record(headers, row)
        if rendered:
            records.append(rendered)
    return "", records


def _extract_xlsx(data: bytes) -> Tuple[str, List[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:  # pragma: no cover
        raise UnsupportedFileType(
            "Excel support requires `openpyxl`. Install via requirements.txt."
        ) from e
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    records: List[str] = []
    for sheet in wb.worksheets:
        rows_iter = sheet.iter_rows(values_only=True)
        rows: List[List] = []
        for row in rows_iter:
            if row is None:
                continue
            if not any(_stringify_cell(c) for c in row):
                continue
            rows.append(list(row))
        if not rows:
            continue
        headers = [_stringify_cell(h) or f"col_{i + 1}" for i, h in enumerate(rows[0])]
        sheet_tag = f"sheet: {sheet.title}"
        for row in rows[1:]:
            rendered = _render_record(headers, row)
            if rendered:
                records.append(f"{sheet_tag}\n{rendered}")
    try:
        wb.close()
    except Exception:  # noqa: BLE001
        pass
    return "", records


def _extract_json(data: bytes) -> Tuple[str, List[str], bool]:
    """Returns (text, records, is_structured)."""
    try:
        payload = json.loads(data.decode("utf-8"))
    except Exception as e:
        raise UnsupportedFileType(f"Invalid JSON: {e}") from e

    def render_obj(obj) -> str:
        if isinstance(obj, str):
            return obj.strip()
        if isinstance(obj, dict):
            lines: List[str] = []
            for k, v in obj.items():
                if v is None or v == "":
                    continue
                if isinstance(v, (dict, list)):
                    lines.append(f"{k}: {json.dumps(v, ensure_ascii=False)}")
                else:
                    lines.append(f"{k}: {v}")
            return "\n".join(lines)
        return json.dumps(obj, ensure_ascii=False)

    if isinstance(payload, list):
        records = [render_obj(item) for item in payload if item is not None]
        records = [r for r in records if r]
        return "", records, True
    if isinstance(payload, dict):
        return render_obj(payload), [], False
    return _stringify_cell(payload), [], False


# ---------------------------------------------------------------------- dispatcher

@dataclass
class Parsed:
    text: str
    records: List[str]
    fmt: str          # FORMAT_LABEL value
    structure: str    # "structured" | "unstructured" | "mixed"


def parse_file(filename: str, data: bytes) -> Parsed:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileType(
            f"Unsupported file type: '{ext}'. Allowed: {sorted(SUPPORTED_EXTENSIONS)}"
        )
    fmt = FORMAT_LABEL[ext]

    if ext in {".txt", ".md", ".markdown"}:
        text, records = _extract_plain_text(data)
        return Parsed(text=text, records=records, fmt=fmt, structure="unstructured")
    if ext == ".pdf":
        text, records = _extract_pdf(data)
        return Parsed(text=text, records=records, fmt=fmt, structure="unstructured")
    if ext in {".html", ".htm"}:
        text, records = _extract_html(data)
        return Parsed(text=text, records=records, fmt=fmt, structure="unstructured")
    if ext == ".docx":
        text, records = _extract_docx(data)
        structure = (
            "mixed" if text and records else "structured" if records else "unstructured"
        )
        return Parsed(text=text, records=records, fmt=fmt, structure=structure)
    if ext == ".csv":
        text, records = _extract_csv(data)
        return Parsed(text=text, records=records, fmt=fmt, structure="structured")
    if ext == ".xlsx":
        text, records = _extract_xlsx(data)
        return Parsed(text=text, records=records, fmt=fmt, structure="structured")
    if ext == ".json":
        text, records, is_structured = _extract_json(data)
        return Parsed(
            text=text,
            records=records,
            fmt=fmt,
            structure="structured" if is_structured else "unstructured",
        )
    raise UnsupportedFileType(f"Unsupported file type: '{ext}'")


# ---------------------------------------------------------------------- store

class DocumentStore:
    """Persists upload metadata + chunk records to a JSON file on disk."""

    def __init__(self, data_dir: Path, chunk_size: int = 600, overlap: int = 100):
        self.data_dir = Path(data_dir)
        self.uploads_dir = self.data_dir / "uploads"
        self.index_path = self.data_dir / "uploaded_chunks.json"
        self.chunk_size = chunk_size
        self.overlap = overlap

        self.data_dir.mkdir(parents=True, exist_ok=True)
        self.uploads_dir.mkdir(parents=True, exist_ok=True)

        self.uploads: List[UploadRecord] = []
        self.chunks: List[ChunkRecord] = []
        self._load()

    # ------------------------------------------------------------------ load/save

    def _load(self) -> None:
        if not self.index_path.exists():
            return
        try:
            with open(self.index_path, "r", encoding="utf-8") as f:
                payload = json.load(f)
            self.uploads = [
                UploadRecord(
                    id=u["id"],
                    filename=u["filename"],
                    category=u["category"],
                    size_bytes=u["size_bytes"],
                    chunks=u["chunks"],
                    chars=u["chars"],
                    uploaded_at=u["uploaded_at"],
                    format=u.get("format", "text"),
                    structure=u.get("structure", "unstructured"),
                    chunk_ids=u.get("chunk_ids", []),
                )
                for u in payload.get("uploads", [])
            ]
            self.chunks = [ChunkRecord(**c) for c in payload.get("chunks", [])]
        except Exception as e:  # noqa: BLE001 — corrupted file, reset
            print(f"[document_store] load failed, resetting: {e}")
            self.uploads = []
            self.chunks = []

    def _save(self) -> None:
        payload = {
            "uploads": [asdict(u) for u in self.uploads],
            "chunks": [asdict(c) for c in self.chunks],
        }
        tmp = self.index_path.with_suffix(".json.tmp")
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False, indent=2)
        tmp.replace(self.index_path)

    # -------------------------------------------------------------- chunk builder

    def _build_chunks(self, parsed: Parsed) -> List[Chunk]:
        prose_chunks: List[Chunk] = []
        record_chunks: List[Chunk] = []

        if parsed.text:
            prose_chunks = chunk_text(
                parsed.text, chunk_size=self.chunk_size, overlap=self.overlap
            )
        if parsed.records:
            header = ""
            # For Excel/CSV-like data, the per-record string itself carries
            # field names; no global header is necessary. Records-only structures
            # are kept atomic by chunk_records.
            record_chunks = chunk_records(
                parsed.records,
                chunk_size=max(self.chunk_size, 800),
                header=header,
            )

        # Concatenate, reindex so chunk_index is monotonic.
        merged: List[Chunk] = []
        idx = 0
        for c in prose_chunks + record_chunks:
            merged.append(Chunk(index=idx, text=c.text))
            idx += 1
        return merged

    # ------------------------------------------------------------------ CRUD

    def add_file(
        self,
        filename: str,
        data: bytes,
        category: Optional[str] = None,
    ) -> Tuple[UploadRecord, List[ChunkRecord]]:
        parsed = parse_file(filename, data)
        if not parsed.text and not parsed.records:
            raise EmptyDocument("No extractable content in the uploaded file.")

        chunks = self._build_chunks(parsed)
        if not chunks:
            raise EmptyDocument("File contained no chunkable text.")

        upload_id = uuid.uuid4().hex[:12]
        category = (category or "uploaded").strip() or "uploaded"
        base = Path(filename).stem
        slug = _slugify(base)
        total = len(chunks)
        chunk_records_list: List[ChunkRecord] = []
        for c in chunks:
            cid = f"up-{slug}-{upload_id}-{c.index:03d}"
            chunk_records_list.append(
                ChunkRecord(
                    id=cid,
                    upload_id=upload_id,
                    title=f"{base} · part {c.index + 1}/{total}",
                    category=category,
                    content=c.text,
                    chunk_index=c.index,
                    total_chunks=total,
                )
            )

        safe_name = f"{upload_id}_{slug}{Path(filename).suffix.lower()}"
        try:
            (self.uploads_dir / safe_name).write_bytes(data)
        except Exception as e:  # noqa: BLE001
            print(f"[document_store] could not persist raw file: {e}")

        record = UploadRecord(
            id=upload_id,
            filename=filename,
            category=category,
            size_bytes=len(data),
            chunks=total,
            chars=sum(len(c.content) for c in chunk_records_list),
            uploaded_at=time.time(),
            format=parsed.fmt,
            structure=parsed.structure,
            chunk_ids=[c.id for c in chunk_records_list],
        )
        self.uploads.append(record)
        self.chunks.extend(chunk_records_list)
        self._save()
        return record, chunk_records_list

    def delete(self, upload_id: str) -> Optional[UploadRecord]:
        record = next((u for u in self.uploads if u.id == upload_id), None)
        if not record:
            return None
        self.uploads = [u for u in self.uploads if u.id != upload_id]
        self.chunks = [c for c in self.chunks if c.upload_id != upload_id]
        for path in self.uploads_dir.glob(f"{upload_id}_*"):
            try:
                path.unlink()
            except OSError:
                pass
        self._save()
        return record

    def list_uploads(self) -> List[dict]:
        return [
            u.to_public()
            for u in sorted(self.uploads, key=lambda u: u.uploaded_at, reverse=True)
        ]
